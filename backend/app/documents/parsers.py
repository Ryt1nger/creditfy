"""Извлечение финансовых показателей из документов: текстовые PDF, XLSX, CSV. Модель — только запасной вариант.

Каждое найденное значение возвращается с подписью источника («Отчёт о финансовых результатах, строка 2110»),
чтобы клиент на шаге «Проверка данных» видел, откуда взялось число, и мог его поправить.
"""
import csv
import io
import re
from dataclasses import dataclass

FIELD_LABELS = {
    "revenue_year": "Выручка за год",
    "net_profit_year": "Чистая прибыль за год",
    "monthly_turnover": "Средние поступления на счёт в месяц",
    "monthly_obligations": "Платежи по кредитам в месяц",
    "purchase_price": "Стоимость покупки по счёту",
}
FIN_FIELDS = ("revenue_year", "net_profit_year", "monthly_turnover", "monthly_obligations")
EXPECTED = {"tax_return": ["revenue_year", "net_profit_year"], "bank_statement": ["monthly_turnover", "monthly_obligations"],
            "invoice": ["purchase_price"]}


@dataclass
class Found:
    field: str
    value: int
    source: str  # parser_pdf | parser_xlsx | parser_csv | ai
    label: str
    confidence: float = 0.95


def _num(s) -> float | None:
    if isinstance(s, (int, float)):
        return float(s)
    if s is None:
        return None
    s = str(s).replace(" ", " ").replace(" ", " ").strip()
    neg = s.startswith("(") and s.endswith(")") or s.startswith("-") or s.startswith("−")
    s = re.sub(r"[^\d,.]", "", s)
    if not s:
        return None
    if s.count(",") == 1 and s.count(".") == 0:
        s = s.replace(",", ".")
    elif s.count(",") and s.count("."):
        s = s.replace(",", "")
    try:
        v = float(s) if s.count(".") <= 1 else float(s.replace(".", ""))
    except ValueError:
        return None
    return -v if neg else v


NUM_RX = r"(\(?[-−]?\d[\d   ]*(?:[.,]\d+)?\)?)"


def pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        return "\n".join((p.extract_text() or "") for p in reader.pages[:10])
    except Exception:  # noqa: BLE001 — битый/сканированный PDF: текста нет
        return ""


def parse_text(text: str, kind: str, src: str, doc_title: str) -> list[Found]:
    out: list[Found] = []
    mult = 1000 if re.search(r"тыс\.?\s*руб", text, re.I) else 1
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]

    def line_value(patterns, field, label):
        for i, ln in enumerate(lines):
            if any(re.search(p, ln, re.I) for p in patterns):
                # значение в той же строке или (при выравнивании по правому краю) в следующей
                for cand in (ln, lines[i + 1] if i + 1 < len(lines) else ""):
                    nums = [n for n in (_num(m) for m in re.findall(NUM_RX, cand)) if n is not None and abs(n) >= 1000 / mult]
                    nums = [n for n in nums if not (1990 <= n <= 2100 or n in (2110, 2120, 2200, 2400))]
                    if nums:
                        out.append(Found(field, int(round(nums[0] * mult)), src, label, 0.9))
                        return

    if kind in ("tax_return", "any"):
        line_value([r"\b2110\b", r"^выручка"], "revenue_year", f"{doc_title}, строка 2110 «Выручка»")
        line_value([r"\b2400\b", r"чистая прибыль"], "net_profit_year", f"{doc_title}, строка 2400 «Чистая прибыль»")
    if kind in ("bank_statement", "any"):
        line_value([r"средн\w*\s+поступлени", r"среднемесячн\w*\s+оборот"], "monthly_turnover", f"{doc_title}, «Среднемесячные поступления»")
        line_value([r"платеж\w*\s+по\s+кредит"], "monthly_obligations", f"{doc_title}, «Платежи по кредитам в месяц»")
    if kind in ("invoice", "any"):
        line_value([r"итого\s+к\s+оплате", r"всего\s+к\s+оплате"], "purchase_price", f"{doc_title}, «Итого к оплате»")
    return out


def parse_rows(rows: list[list], kind: str, src: str, doc_title: str) -> list[Found]:
    """Таблица: выписка (помесячные строки) или отчётность (код строки + значение)."""
    out: list[Found] = []
    rows = [[c for c in r] for r in rows if any(c not in (None, "") for c in r)]
    # отчётность: строки с кодами 2110 / 2400
    for r in rows:
        cells = [str(c).strip() if c is not None else "" for c in r]
        code = next((c for c in cells if c in ("2110", "2400")), None)
        if code:
            nums = [_num(c) for c in r if not (str(c).strip() in ("2110", "2400"))]
            nums = [n for n in nums if n is not None and abs(n) >= 1000]
            if nums:
                field = "revenue_year" if code == "2110" else "net_profit_year"
                name = "Выручка" if code == "2110" else "Чистая прибыль"
                out.append(Found(field, int(round(nums[0])), src, f"{doc_title}, строка {code} «{name}»"))
    # выписка: заголовок с «Поступления» и, возможно, «платежи по кредитам»
    for i, r in enumerate(rows):
        head = [str(c).lower() if c is not None else "" for c in r]
        inc = next((j for j, h in enumerate(head) if "поступлен" in h or h.startswith("приход")), None)
        if inc is None:
            continue
        loan = next((j for j, h in enumerate(head) if "кредит" in h and "платеж" in h), None)
        body = [x for x in rows[i + 1:] if x and not str(x[0] or "").lower().startswith(("итого", "всего", "среднее"))]
        incs = [v for v in (_num(x[inc]) for x in body if inc < len(x)) if v is not None and v > 0]
        if incs:
            out.append(Found("monthly_turnover", int(round(sum(incs) / len(incs))), src,
                             f"{doc_title}: среднее поступлений за {len(incs)} мес."))
        if loan is not None:
            loans = [v for v in (_num(x[loan]) for x in body if loan < len(x)) if v is not None and v >= 0]
            if loans:
                out.append(Found("monthly_obligations", int(round(sum(loans) / len(loans))), src,
                                 f"{doc_title}: среднее платежей по кредитам за {len(loans)} мес."))
        break
    return out


def xlsx_rows(content: bytes) -> list[list]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    rows: list[list] = []
    for ws in wb.worksheets[:3]:
        for r in ws.iter_rows(values_only=True, max_row=500):
            rows.append(list(r))
    return rows


def csv_rows(content: bytes) -> list[list]:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        return []
    try:
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    return [r for r in csv.reader(io.StringIO(text), dialect)][:500]


def extract(kind: str, filename: str, content: bytes) -> tuple[list[Found], str, str]:
    """→ (найденные значения, метод, текст для запасного разбора моделью)."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    title = {"tax_return": "Отчётность", "bank_statement": "Выписка по счёту", "invoice": "Счёт"}.get(kind, "Документ")
    try:
        if ext == "pdf":
            text = pdf_text(content)
            return parse_text(text, kind, "parser_pdf", title), "parser_pdf", text
        if ext == "xlsx":
            rows = xlsx_rows(content)
            return parse_rows(rows, kind, "parser_xlsx", title), "parser_xlsx", "\n".join(";".join("" if c is None else str(c) for c in r) for r in rows[:80])
        if ext == "csv":
            rows = csv_rows(content)
            return parse_rows(rows, kind, "parser_csv", title), "parser_csv", "\n".join(";".join(r) for r in rows[:80])
    except Exception:  # noqa: BLE001 — повреждённый файл не должен ронять загрузку
        return [], "error", ""
    return [], "none", ""
